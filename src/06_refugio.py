"""Módulo F — "Refugio austral": Argentina frente a Nueva Zelanda, Uruguay, Chile, Portugal y Canadá.

Compara dimensiones a favor (paz, alimentos, energía, litio, patrimonio natural) y CONTRAPESOS (defaults, inflación,
riesgo país, calificación soberana, controles de capital, Estado de derecho) sin agregarlas en un índice compuesto.
Riesgo país: el EMBI (JP Morgan) es propietario; se usa la cifra que el BCRA cita textualmente en el IPOM (solo
Argentina) y, como comparable para los seis, la calificación soberana publicada por cada gobierno/oficina de deuda.

Salidas:
  data/processed/F_*.csv          series limpias y tabla resumen (un DATO/ESTIMACIÓN por celda, con año)
  outputs/charts/F_*.{png,svg}    small multiples por dimensión, Vaca Muerta, inflación, defaults, WGI
  docs/claims/claims_F.csv        afirmaciones con cita literal o localizador reproducible

Corre de punta a punta desde las copias de data/raw (descarga solo lo que falte).
"""
from __future__ import annotations

import csv
import gzip
import json
import re
import shutil
import zipfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402

from common import CHARTS, PROCESSED, RAW, download, local_text, quote, sha256, write_ledger  # noqa: E402

MODULO = "F"
ISO3 = ["ARG", "NZL", "URY", "CHL", "PRT", "CAN"]
NOMBRE = {"ARG": "Argentina", "NZL": "Nueva Zelanda", "URY": "Uruguay", "CHL": "Chile", "PRT": "Portugal",
          "CAN": "Canadá"}
EN = {"ARG": "Argentina", "NZL": "New Zealand", "URY": "Uruguay", "CHL": "Chile", "PRT": "Portugal", "CAN": "Canada"}
ISO2 = {"ARG": "ar", "NZL": "nz", "URY": "uy", "CHL": "cl", "PRT": "pt", "CAN": "ca"}
ROOT = RAW.parents[1]
FECHA_CORTE = "2026-10-03"  # fecha de corte del módulo (para afirmaciones de cálculo propio sin archivo local)

D360 = "https://data360api.worldbank.org/data360/data"
CKAN_RES = "https://datos.energia.gob.ar/dataset/c846e79c-026c-4040-897f-1ad3543b407c/resource/"
WGI = {"GOV_WGI_RL": "Rule of Law", "GOV_WGI_PV": "Political Stability and Absence of Violence",
       "GOV_WGI_GE": "Government Effectiveness", "GOV_WGI_CC": "Control of Corruption"}

SRC = {
    # id: (url, nombre, ext, descripción, tipo)
    **{f"wgi_{k[-2:]}": (f"{D360}?DATABASE_ID=WB_WGI&INDICATOR={k}&REF_AREA={','.join(ISO3)}"
                         "&COMP_BREAKDOWN_1=WGI_EST,WGI_SC",
                         f"F_wb360_wgi_{k[-2:].lower()}", "json",
                         f"Banco Mundial, Worldwide Governance Indicators (API Data360), {v}", "P")
       for k, v in WGI.items()},
    "wdi_cpi": (f"{D360}?DATABASE_ID=WB_WDI&INDICATOR=WB_WDI_FP_CPI_TOTL_ZG&REF_AREA={','.join(ISO3)}",
                "F_wb360_wdi_fp_cpi_totl_zg", "json", "Banco Mundial, WDI FP.CPI.TOTL.ZG (inflación, IPC, % anual)", "P"),
    "wdi_eimp": (f"{D360}?DATABASE_ID=WB_WDI&INDICATOR=WB_WDI_EG_IMP_CONS_ZS&REF_AREA={','.join(ISO3)}",
                 "F_wb360_wdi_eg_imp_cons_zs", "json",
                 "Banco Mundial, WDI EG.IMP.CONS.ZS (importaciones netas de energía, % del uso)", "P"),
    "gpi": ("https://www.visionofhumanity.org/wp-content/uploads/2026/06/Global-Peace-Index-2026-Report.pdf",
            "F_iep_gpi2026_report", "pdf", "Institute for Economics & Peace, Global Peace Index 2026 (informe)", "R"),
    "wjp": ("https://worldjusticeproject.org/rule-of-law-index/downloads/WJPIndex2025.pdf",
            "F_wjp_index2025", "pdf", "World Justice Project, Rule of Law Index 2025 (informe)", "R"),
    "fao": ("https://bulks-faostat.fao.org/production/FoodBalanceSheets_E_All_Data_(Normalized).zip",
            "F_faostat_fbs_normalized", "zip", "FAOSTAT, Food Balance Sheets (2010-), descarga masiva normalizada", "P"),
    "oil": (CKAN_RES + "af8c50bb-fde0-43b7-98eb-7cd14daf586c/download/"
            "serie-histrica-de-produccin-de-petrleo-por-cuenca-y-sub-tipo-de-recurso-captulo-iv-.csv",
            "F_energia_serie_petroleo_cuenca_subtipo", "csv",
            "Secretaría de Energía, producción de petróleo por cuenca y subtipo de recurso (Capítulo IV)", "P"),
    "gas": (CKAN_RES + "a3244ddd-38bc-4800-a700-360b649d2f3a/download/"
            "serie-histrica-de-produccin-de-gas-natural-por-cuenca-y-sub-tipo-de-recurso-captulo-iv-.csv",
            "F_energia_serie_gas_cuenca_subtipo", "csv",
            "Secretaría de Energía, producción de gas natural por cuenca y subtipo de recurso (Capítulo IV)", "P"),
    "vm": (CKAN_RES + "b5b58cdc-9e07-41f9-b392-fb9ec68b0725/download/"
           "produccin-de-pozos-de-gas-y-petrleo-no-convencional.csv",
           "F_energia_pozos_no_convencional", "csv.gz",
           "Secretaría de Energía, producción por pozo no convencional (Capítulo IV); copia comprimida con gzip", "P"),
    "ben": ("http://www.energia.gob.ar/contenidos/archivos/Reorganizacion/informacion_del_mercado/publicaciones/"
            "energia_en_gral/balances_2025/balance_2025_v0_h.xlsx",
            "F_energia_ben2025", "xlsx", "Secretaría de Energía, Balance Energético Nacional 2025 (revisión 0, provisorio)", "P"),
    "usgs": ("https://pubs.usgs.gov/periodicals/mcs2026/mcs2026-lithium.pdf", "F_usgs_mcs2026_lithium", "pdf",
             "USGS, Mineral Commodity Summaries 2026: Lithium", "P"),
    "unesco": ("https://data.unesco.org/api/explore/v2.1/catalog/datasets/whc001/exports/csv?select=id_no,name_en,"
               "category,states_names,iso_codes,date_inscribed,danger,transboundary&delimiter=%3B",
               "F_unesco_whc001", "csv", "UNESCO Open Data, World Heritage List (dataset whc001)", "P"),
    "boc": ("https://www.bankofcanada.ca/wp-content/uploads/2025/10/BoC-BoE-Database-2025.xlsx",
            "F_boc_boe_sovereign_default_db2025", "xlsx",
            "Bank of Canada–Bank of England, Sovereign Default Database 2025 (datos 1960–2024)", "P"),
    "boc_note": ("https://www.bankofcanada.ca/wp-content/uploads/2025/10/san2025-24.pdf", "F_boc_boe_san2025-24", "pdf",
                 "Bank of Canada, Staff Analytical Note 2025-24 (metodología de la base de defaults)", "P"),
    "ipc": ("https://apis.datos.gob.ar/series/api/series/?ids=148.3_INIVELNAL_DICI_M_26,"
            "148.3_INIVELNAL_DICI_M_26:percent_change,148.3_INIVELNAL_DICI_M_26:percent_change_a_year_ago"
            "&format=csv&limit=1000",
            "F_datosgob_ipc_nacional", "csv", "INDEC IPC Nacional (base dic-2016) vía API de Series de Tiempo de datos.gob.ar", "P"),
    "a5245": ("https://www.bcra.gob.ar/archivos/Pdfs/comytexord/A5245.pdf", "F_bcra_comA5245", "pdf",
              "BCRA, Comunicación \"A\" 5245 (10/11/2011)", "P"),
    "a5850": ("https://www.bcra.gob.ar/archivos/Pdfs/comytexord/A5850.pdf", "F_bcra_comA5850", "pdf",
              "BCRA, Comunicación \"A\" 5850 (17/12/2015)", "P"),
    "dnu609": ("https://servicios.infoleg.gob.ar/infolegInternet/anexos/325000-329999/327566/norma.htm",
               "F_infoleg_dnu609_2019", "html", "InfoLEG, Decreto DNU 609/2019 (01/09/2019)", "P"),
    "a6770": ("https://www.bcra.gob.ar/archivos/Pdfs/comytexord/A6770.pdf", "F_bcra_comA6770", "pdf",
              "BCRA, Comunicación \"A\" 6770 (01/09/2019)", "P"),
    "a6815": ("https://www.bcra.gob.ar/archivos/Pdfs/comytexord/A6815.pdf", "F_bcra_comA6815", "pdf",
              "BCRA, Comunicación \"A\" 6815 (28/10/2019)", "P"),
    "a8226": ("https://www.bcra.gob.ar/archivos/Pdfs/comytexord/A8226.pdf", "F_bcra_comA8226", "pdf",
              "BCRA, Comunicación \"A\" 8226 (11/04/2025)", "P"),
    "dec1570": ("https://servicios.infoleg.gob.ar/infolegInternet/anexos/70000-74999/70355/norma.htm",
                "F_infoleg_dec1570_2001", "html", "InfoLEG, Decreto 1570/2001 (01/12/2001, \"corralito\")", "P"),
    "cancilleria": ("https://www.cancilleria.gob.ar/es/actualidad/noticias/"
                    "argentina-y-estados-unidos-firmaron-un-acuerdo-sobre-comercio-e-inversiones",
                    "F_cancilleria_acuerdo_eeuu", "html",
                    "Cancillería argentina, nota del 05/02/2026 sobre el Acuerdo de Comercio e Inversiones Recíprocos", "P"),
    "esf": ("https://home.treasury.gov/system/files/206/ESF-October-2025-FS_Trunc_Notes.pdf",
            "F_treasury_esf_oct2025", "pdf", "U.S. Treasury, Exchange Stabilization Fund, estados de octubre de 2025", "P"),
    # --- pendientes cerrados el 2026-10-03 (riesgo país, calificaciones, controles de capital vigentes, MNNA)
    "ipom": ("https://www.bcra.gob.ar/archivos/Pdfs/PublicacionesEstadisticas/informes/informe-politica-monetaria-2026-T2.pdf",
             "F_bcra_ipom_2026t2", "pdf",
             "BCRA, Informe de Política Monetaria, segundo trimestre de 2026 (publicado el 06/08/2026)", "P"),
    "texord": ("https://www.bcra.gob.ar/archivos/Pdfs/Texord/t-excbio.pdf", "F_bcra_texord_exterior_cambios", "pdf",
               "BCRA, Texto ordenado \"Exterior y Cambios\" al 14/09/2026 (última comunicación incorporada: A 8481)", "P"),
    "bcra_nec": ("https://www.bcra.gob.ar/normativa-de-exterior-y-cambios/", "F_bcra_normativa_exterior_cambios", "html",
                 "BCRA, página \"Normativa de Exterior y Cambios\" (puntos principales de la normativa vigente)", "P"),
    "oecd_lt": ("https://sdmx.oecd.org/public/rest/data/OECD.SDD.STES,DSD_STES@DF_FINMARK,/"
                "ARG+URY+NZL+CHL+PRT+CAN.M.IRLT......?startPeriod=2024-01&format=csvfile",
                "F_oecd_irlt_mensual", "csv",
                "OCDE, Financial market (DF_FINMARK), Long-term interest rates (IRLT, bonos del gobierno a ~10 años, % anual)", "P"),
    "rt_nz": ("https://web.archive.org/web/20260409193619id_/https://debtmanagement.treasury.govt.nz/investor-resources/credit-ratings",
              "F_nzdm_credit_ratings_wayback20260409", "html",
              "New Zealand Debt Management (Treasury), \"Credit ratings\" (captura Wayback del 09/04/2026)", "P"),
    "rt_uy": ("https://deuda.mef.gub.uy/6475/14/areas/calificacion-crediticia.html", "F_mef_uy_calificacion_crediticia", "html",
              "MEF Uruguay, Unidad de Gestión de Deuda, \"Calificación Crediticia\"", "P"),
    "rt_cl": ("https://www.hacienda.cl/areas-de-trabajo/finanzas-internacionales/oficina-de-la-deuda-publica/estadisticas/"
              "ratings-historicos", "F_hacienda_cl_ratings_historicos", "html",
              "Ministerio de Hacienda de Chile, Oficina de la Deuda Pública, \"Ratings históricos\"", "P"),
    "rt_pt": ("https://www.igcp.pt/sites/default/files/2026-09/IGCP_Investor_Presentation.pdf", "F_igcp_investor_presentation_2026-09",
              "pdf", "IGCP (agencia de deuda de Portugal), Investor Presentation, septiembre de 2026", "P"),
    "rt_ca": ("https://www.canada.ca/en/department-finance/services/publications/debt-management-report/2024-2025.html",
              "F_finance_canada_dmr_2024-25", "html", "Department of Finance Canada, Debt Management Report 2024–25", "P"),
    "mnna": ("https://web.archive.org/web/20260924124444id_/https://www.state.gov/major-non-nato-ally-status",
             "F_state_mnna_wayback20260924", "html",
             "U.S. Department of State, \"Major Non-NATO Ally Status\" (captura Wayback del 24/09/2026)", "P"),
}

# Escala de calificaciones: escalones por debajo de AAA/Aaa (S&P y Fitch | Moody's)
NOTCH_SP = ["AAA", "AA+", "AA", "AA-", "A+", "A", "A-", "BBB+", "BBB", "BBB-", "BB+", "BB", "BB-", "B+", "B", "B-",
            "CCC+", "CCC", "CCC-"]
NOTCH_MD = ["Aaa", "Aa1", "Aa2", "Aa3", "A1", "A2", "A3", "Baa1", "Baa2", "Baa3", "Ba1", "Ba2", "Ba3", "B1", "B2", "B3",
            "Caa1", "Caa2", "Caa3"]

FILES: dict[str, Path] = {}
TEXTS: dict[str, str] = {}
CLAIMS: list[dict] = []


# ------------------------------------------------------------------------------------------------ utilidades
def fetch(sid: str) -> Path:
    url, name, ext, _d, _t = SRC[sid]
    if ext == "csv.gz":  # archivo grande (>100 MB): se guarda comprimido, mismo contenido byte a byte
        found = sorted(RAW.glob(f"{name}_*.csv.gz"))
        if found:
            return found[-1]
        p = download(url, name, "csv", timeout=1800)
        gz = p.with_suffix(".csv.gz")
        with p.open("rb") as fi, gzip.open(gz, "wb", compresslevel=9) as fo:
            shutil.copyfileobj(fi, fo)
        p.unlink()
        return gz
    return download(url, name, ext, timeout=600)


def text(sid: str) -> str:
    """Texto normalizado de un PDF/HTML. Para PDFs se cachea junto a la copia (data/raw/*.txt)."""
    if sid in TEXTS:
        return TEXTS[sid]
    p = FILES[sid]
    if p.suffix == ".pdf":
        cache = p.with_suffix(".txt")
        if not cache.exists():
            cache.write_text(local_text(p), encoding="utf-8")
        TEXTS[sid] = cache.read_text(encoding="utf-8")
    else:
        TEXTS[sid] = local_text(p)
    return TEXTS[sid]


def rel(p: Path) -> str:
    return p.relative_to(ROOT).as_posix()


def claim(cid, etiqueta, afirmacion, valor, sid, cita=None, needle=None, needles=None):
    url, _n, _e, desc, tipo = SRC[sid] if sid else ("", "", "", "Cálculo propio (Colossus Lab)", "")
    if needle is not None:
        cita = quote(text(sid), needle)
    if needles is not None:  # varias citas literales del mismo documento, separadas por " | " (formato de 09_auditoria)
        cita = " | ".join(quote(text(sid), n) for n in needles)
    f = FILES.get(sid)
    # fecha de acceso = fecha de la copia local (sufijo _YYYY-MM-DD); cálculos propios: fecha de corte del módulo
    m = re.search(r"_(\d{4}-\d{2}-\d{2})\.", f.name) if f else None
    CLAIMS.append(dict(claim_id=cid, etiqueta=etiqueta, afirmacion=afirmacion, valor=valor, fuente=desc,
                       tipo_fuente=tipo, url=url, archivo_local=rel(f) if f else "",
                       sha256=sha256(f) if f else "", cita_textual=cita,
                       fecha_acceso=m.group(1) if m else FECHA_CORTE))
    return cid


_n = 0


def nid() -> str:
    global _n
    _n += 1
    return f"F{_n:03d}"


# ------------------------------------------------------------------------------------------------ datos
def load_d360(sid: str) -> pd.DataFrame:
    js = json.loads(FILES[sid].read_text(encoding="utf-8"))
    df = pd.DataFrame(js["value"])
    df["OBS_VALUE"] = pd.to_numeric(df["OBS_VALUE"], errors="coerce")
    df["TIME_PERIOD"] = df["TIME_PERIOD"].astype(int)
    return df


def fao_subset() -> pd.DataFrame:
    """Filtra la hoja de balance FAOSTAT a los 6 países (cacheado en data/processed)."""
    out = PROCESSED / "F_fao_fbs_6paises.csv"
    if out.exists() and out.stat().st_mtime >= FILES["fao"].stat().st_mtime:
        return pd.read_csv(out)
    z = zipfile.ZipFile(FILES["fao"])
    parts = []
    for ch in pd.read_csv(z.open("FoodBalanceSheets_E_All_Data_(Normalized).csv"), encoding="latin-1",
                          chunksize=2_000_000, dtype={"Area Code (M49)": str, "Item Code (FBS)": str}):
        parts.append(ch[ch["Area"].isin(EN.values())])
    df = pd.concat(parts)
    df.to_csv(out, index=False)
    return df


FAO_GROUPS = {2905: "Cereales", 2907: "Raíces amiláceas", 2909: "Azúcar y endulzantes", 2911: "Legumbres",
              2912: "Frutos secos", 2913: "Oleaginosas", 2914: "Aceites vegetales", 2918: "Hortalizas",
              2919: "Frutas", 2943: "Carne", 2945: "Menudencias", 2946: "Grasas animales", 2948: "Leche",
              2949: "Huevos", 2960: "Pescados y mariscos"}
FAO_YEARS = list(range(2019, 2024))


def main() -> None:
    for sid in SRC:
        FILES[sid] = fetch(sid)
    PROCESSED.mkdir(parents=True, exist_ok=True)
    CHARTS.mkdir(parents=True, exist_ok=True)

    table: list[dict] = []  # tabla resumen: una fila por (indicador, país)

    def cell(dim, ind, iso, valor, anio, unidad, etiqueta, cid, mejor="", nota=""):
        table.append(dict(dimension=dim, indicador=ind, pais=iso, valor=valor, anio=anio, unidad=unidad,
                          etiqueta=etiqueta, claim_id=cid, mejor=mejor, nota=nota))

    # ===================================================================== 1. Paz y gobernanza
    wgi_frames = []
    for k, name in WGI.items():
        sid = f"wgi_{k[-2:]}"
        df = load_d360(sid)
        df = df[df.INDICATOR == k]
        wgi_frames.append(df.assign(indicador=name))
        for iso in ISO3:
            s = df[(df.REF_AREA == iso) & (df.COMP_BREAKDOWN_1 == "WGI_EST")].sort_values("TIME_PERIOD")
            last = s.iloc[-1]
            cid = claim(nid(), "DATO", f"WGI {name}: estimación de {NOMBRE[iso]} (escala aprox. −2,5 a +2,5)",
                        f"{last.OBS_VALUE:.2f} ({last.TIME_PERIOD})", sid,
                        cita=f"archivo JSON Data360; INDICATOR={k}; REF_AREA={iso}; COMP_BREAKDOWN_1=WGI_EST; "
                             f"TIME_PERIOD={last.TIME_PERIOD}; OBS_VALUE={last.OBS_VALUE}")
            cell("1 Paz y gobernanza", f"WGI {name} (estimación)", iso, round(last.OBS_VALUE, 2),
                 int(last.TIME_PERIOD), "−2,5 a +2,5", "DATO", cid, "mayor")
    wgi = pd.concat(wgi_frames)
    wgi_est = (wgi[wgi.COMP_BREAKDOWN_1 == "WGI_EST"][["indicador", "REF_AREA", "TIME_PERIOD", "OBS_VALUE"]]
               .rename(columns={"REF_AREA": "pais", "TIME_PERIOD": "anio", "OBS_VALUE": "estimacion"})
               .sort_values(["indicador", "pais", "anio"]))
    wgi_est.to_csv(PROCESSED / "F_wgi_serie.csv", index=False)

    # Global Peace Index 2026 (puesto y puntaje; 1 = más pacífico)
    gpi = {"ARG": (72, "1.922"), "NZL": (2, "1.343"), "URY": (43, "1.754"), "CHL": (52, "1.826"),
           "PRT": (7, "1.427"), "CAN": (14, "1.525")}
    for iso, (rk, sc) in gpi.items():
        cid = claim(nid(), "DATO", f"Global Peace Index 2026: puesto y puntaje de {NOMBRE[iso]} (1 = más pacífico; 163 países)",
                    f"puesto {rk}, puntaje {sc}", "gpi", needle=f"{rk} {EN[iso]} {sc}")
        cell("1 Paz y gobernanza", "Global Peace Index: puntaje (menor = más pacífico)", iso, float(sc), 2026,
             "1 a 5", "DATO", cid, "menor", f"puesto {rk} de 163")
    gpi_n = claim(nid(), "DATO", "El GPI 2026 cubre 163 países", "163", "gpi", needle="Of the 163 countries on the Index")
    gpi_arg = claim(nid(), "DATO", "Argentina tuvo el mayor deterioro porcentual de la región en el GPI 2026 (−6,1%), "
                    "por el dominio de conflicto en curso", "6,1%", "gpi",
                    needle="Argentina recorded the largest percentage deterioration in the")
    claim(nid(), "DATO", "Detalle: el puntaje de Argentina empeoró 6,1% en el GPI 2026", "6,1%", "gpi",
          needle="region on the 2026 GPI, with its overall score deteriorating by 6.1 per cent")

    # ===================================================================== 2. Autosuficiencia alimentaria (FAOSTAT)
    fbs = fao_subset()
    fb = fbs[fbs["Item Code"].isin(FAO_GROUPS) & fbs.Year.isin(FAO_YEARS)]
    prod = fb[fb["Element Code"] == 5511].groupby(["Area", "Item Code"]).Value.sum()
    supl = fb[fb["Element Code"] == 5301].groupby(["Area", "Item Code"]).Value.sum()
    kcal = fb[fb["Element Code"] == 664].groupby(["Area", "Item Code"]).Value.mean()
    ssr = (prod / supl * 100).rename("ssr_pct")
    fao = pd.concat([prod.rename("produccion_kt_2019_23"), supl.rename("suministro_interno_kt_2019_23"), ssr,
                     kcal.rename("kcal_cap_dia_prom")], axis=1).reset_index()
    fao["grupo"] = fao["Item Code"].map(FAO_GROUPS)
    fao["pais"] = fao.Area.map({v: k for k, v in EN.items()})
    # 2023 sola, para mostrar el efecto de la sequía argentina 2022/23
    f23 = fbs[fbs["Item Code"].isin(FAO_GROUPS) & (fbs.Year == 2023)]
    s23 = (f23[f23["Element Code"] == 5511].groupby(["Area", "Item Code"]).Value.sum()
           / f23[f23["Element Code"] == 5301].groupby(["Area", "Item Code"]).Value.sum() * 100)
    fao = fao.merge(s23.rename("ssr_pct_2023").reset_index(), on=["Area", "Item Code"], how="left")
    fao[["pais", "grupo", "Item Code", "produccion_kt_2019_23", "suministro_interno_kt_2019_23", "ssr_pct",
         "ssr_pct_2023", "kcal_cap_dia_prom"]].round(2).to_csv(PROCESSED / "F_fao_autosuficiencia.csv", index=False)

    # insumos de Argentina (DATO) para trazar la estimación de cereales
    a = fbs[(fbs.Area == "Argentina") & (fbs["Item Code"] == 2905) & (fbs.Year == 2023)]
    p23 = a[a["Element Code"] == 5511].Value.iloc[0]
    d23 = a[a["Element Code"] == 5301].Value.iloc[0]
    c_p = claim(nid(), "DATO", "FAOSTAT: producción de cereales (excl. cerveza) de Argentina, 2023", f"{p23:,.0f} miles de t",
                "fao", cita=f"Area=Argentina; Item Code=2905; Element Code=5511; Year=2023; Value={p23}")
    c_d = claim(nid(), "DATO", "FAOSTAT: suministro interno de cereales de Argentina, 2023", f"{d23:,.0f} miles de t",
                "fao", cita=f"Area=Argentina; Item Code=2905; Element Code=5301; Year=2023; Value={d23}")
    for iso in ISO3:
        r = fao[(fao.pais == iso) & (fao["Item Code"] == 2905)].iloc[0]
        cid = claim(nid(), "ESTIMACIÓN", f"Autosuficiencia en cereales de {NOMBRE[iso]}: producción / suministro interno, "
                    "suma 2019–2023", f"{r.ssr_pct:.0f}%", "fao",
                    cita=f"SSR = Σ Producción(5511) / Σ Suministro interno(5301), Item 2905, Area={EN[iso]}, "
                         f"años 2019–2023 = {r.produccion_kt_2019_23:.0f}/{r.suministro_interno_kt_2019_23:.0f}"
                         + (f"; insumos 2023: {c_p}, {c_d}" if iso == "ARG" else ""))
        cell("2 Alimentos", "Autosuficiencia en cereales (prod./suministro interno)", iso, round(r.ssr_pct), "2019–2023",
             "%", "ESTIMACIÓN", cid, "mayor", f"2023: {r.ssr_pct_2023:.0f}%")
        g = fao[fao.pais == iso]
        cov = (g.kcal_cap_dia_prom * g.ssr_pct.clip(upper=100)).sum() / g.kcal_cap_dia_prom.sum()
        cid = claim(nid(), "ESTIMACIÓN", f"Cobertura calórica doméstica de {NOMBRE[iso]}: Σ_g kcal_g × min(SSR_g,100%) / Σ kcal_g "
                    "sobre 15 grupos FAOSTAT, 2019–2023", f"{cov:.0f}%", "fao",
                    cita="kcal_g = promedio 2019–2023 de Food supply (kcal/capita/day, Element 664) por grupo; "
                         "SSR_g = Σ Producción(5511)/Σ Suministro interno(5301) 2019–2023; grupos Item Code "
                         + ",".join(str(k) for k in FAO_GROUPS) + f"; Area={EN[iso]}")
        cell("2 Alimentos", "Cobertura calórica doméstica (kcal ponderadas, SSR tope 100%)", iso, round(cov), "2019–2023",
             "%", "ESTIMACIÓN", cid, "mayor")

    # ===================================================================== 3. Energía
    eimp = load_d360("wdi_eimp")
    for iso in ISO3:
        s = eimp[(eimp.REF_AREA == iso) & eimp.OBS_VALUE.notna()].sort_values("TIME_PERIOD").iloc[-1]
        cid = claim(nid(), "DATO", f"Importaciones netas de energía de {NOMBRE[iso]} (% del uso de energía; negativo = exportador neto)",
                    f"{s.OBS_VALUE:.1f}% ({s.TIME_PERIOD})", "wdi_eimp",
                    cita=f"INDICATOR=WB_WDI_EG_IMP_CONS_ZS; REF_AREA={iso}; TIME_PERIOD={s.TIME_PERIOD}; OBS_VALUE={s.OBS_VALUE}")
        cell("3 Energía", "Importaciones netas de energía (% del uso)", iso, round(s.OBS_VALUE, 1), int(s.TIME_PERIOD),
             "%", "DATO", cid, "menor")
    (eimp[["REF_AREA", "TIME_PERIOD", "OBS_VALUE"]].rename(columns={"REF_AREA": "pais", "TIME_PERIOD": "anio",
                                                                     "OBS_VALUE": "imp_netas_energia_pct_uso"})
     .sort_values(["pais", "anio"]).to_csv(PROCESSED / "F_energia_importaciones_wb.csv", index=False))

    # Balance Energético Nacional 2025 (ktep): misma definición que el WB, con datos propios de la SE
    ws = openpyxl.load_workbook(FILES["ben"], data_only=True)["FINAL HOR"]
    row_tot = next(r for r in ws.iter_rows(values_only=True) if "TOTAL I" in r)
    vals_tot = [x for x in row_tot[row_tot.index("TOTAL I") + 1:] if x is not None]
    ben_prod, _imp, _stk, ben_exp, _noap, _perd, _aj, ben_oferta = vals_tot[:8]
    c_bp = claim(nid(), "DATO", "BEN 2025 (provisorio): producción de energía primaria", f"{ben_prod:,.0f} ktep", "ben",
                 cita=f"hoja FINAL HOR; fila 'TOTAL I'; columna PRODUCCIÓN; valor={ben_prod}")
    c_bo = claim(nid(), "DATO", "BEN 2025 (provisorio): oferta interna de energía primaria", f"{ben_oferta:,.0f} ktep", "ben",
                 cita=f"hoja FINAL HOR; fila 'TOTAL I'; columna OFERTA INTERNA; valor={ben_oferta}")
    claim(nid(), "DATO", "BEN 2025 (provisorio): exportación y búnker de energía primaria", f"{-ben_exp:,.0f} ktep", "ben",
          cita=f"hoja FINAL HOR; fila 'TOTAL I'; columna EXPORTACIÓN Y BUNKER; valor={ben_exp}")
    imp_ben = (ben_oferta - ben_prod) / ben_oferta * 100
    c_ben = claim(nid(), "ESTIMACIÓN", "Argentina 2025: importaciones netas de energía primaria, % de la oferta interna "
                  "(definición análoga a EG.IMP.CONS.ZS)", f"{imp_ben:.1f}%", "ben",
                  cita=f"(oferta interna − producción)/oferta interna = ({ben_oferta:.0f} − {ben_prod:.0f})/{ben_oferta:.0f}; "
                       f"insumos {c_bp}, {c_bo}")

    # Producción de petróleo y gas: nacional (series) y Vaca Muerta (pozo a pozo)
    oil = pd.read_csv(FILES["oil"], encoding="utf-8-sig")
    gas = pd.read_csv(FILES["gas"], encoding="utf-8-sig")
    vm = pd.read_csv(FILES["vm"], usecols=["anio", "mes", "prod_pet", "prod_gas", "formacion"], low_memory=False)
    vm = vm[vm.formacion.astype(str).str.strip().str.lower() == "vaca muerta"]
    vmm = vm.groupby(["anio", "mes"])[["prod_pet", "prod_gas"]].sum().reset_index()
    vmm["indice_tiempo"] = vmm.apply(lambda r: f"{int(r.anio)}-{int(r.mes):02d}", axis=1)
    vmm["dias"] = pd.to_datetime(vmm.indice_tiempo + "-01").dt.days_in_month
    vmm["vm_petroleo_kbbl_d"] = vmm.prod_pet * 6.28981 / vmm.dias / 1000
    vmm["vm_gas_mmm3_d"] = vmm.prod_gas / vmm.dias / 1000
    en = (oil[["indice_tiempo", "total", "kbbl_diario", "shale"]]
          .merge(gas[["indice_tiempo", "produccion_gas_natural_total", "produccion_gas_natural_total_diario"]],
                 on="indice_tiempo")
          .merge(vmm[["indice_tiempo", "prod_pet", "prod_gas", "vm_petroleo_kbbl_d", "vm_gas_mmm3_d"]],
                 on="indice_tiempo", how="left"))
    en.columns = ["mes", "pet_total_m3", "pet_total_kbbl_d", "pet_shale_m3", "gas_total_miles_m3", "gas_total_mmm3_d",
                  "vm_pet_m3", "vm_gas_miles_m3", "vm_pet_kbbl_d", "vm_gas_mmm3_d"]
    en["vm_share_pet_pct"] = en.vm_pet_m3 / en.pet_total_m3 * 100
    en["vm_share_gas_pct"] = en.vm_gas_miles_m3 / en.gas_total_miles_m3 * 100
    en.round(3).to_csv(PROCESSED / "F_vaca_muerta_mensual.csv", index=False)
    last = en.dropna(subset=["vm_pet_m3"]).iloc[-1]
    yago = en[en.mes == f"{int(last.mes[:4]) - 1}{last.mes[4:]}"].iloc[0]
    lm = last.mes
    c_vmp = claim(nid(), "DATO", f"Producción de petróleo de la formación Vaca Muerta, {lm} (DDJJ pozo a pozo, provisoria)",
                  f"{last.vm_pet_m3:,.0f} m³", "vm",
                  cita=f"formacion='vaca muerta'; anio={lm[:4]}; mes={int(lm[5:])}; Σ prod_pet = {last.vm_pet_m3:.3f}")
    c_vmg = claim(nid(), "DATO", f"Producción de gas de la formación Vaca Muerta, {lm}", f"{last.vm_gas_miles_m3:,.0f} miles de m³",
                  "vm", cita=f"formacion='vaca muerta'; anio={lm[:4]}; mes={int(lm[5:])}; Σ prod_gas = {last.vm_gas_miles_m3:.3f}")
    c_pt = claim(nid(), "DATO", f"Producción total de petróleo de Argentina, {lm}", f"{last.pet_total_kbbl_d:.0f} mil bbl/d",
                 "oil", cita=f"indice_tiempo={lm}; total={last.pet_total_m3}; kbbl_diario={last.pet_total_kbbl_d}")
    c_gt = claim(nid(), "DATO", f"Producción total de gas de Argentina, {lm}", f"{last.gas_total_mmm3_d:.1f} MMm³/d",
                 "gas", cita=f"indice_tiempo={lm}; produccion_gas_natural_total={last.gas_total_miles_m3}; "
                            f"produccion_gas_natural_total_diario={last.gas_total_mmm3_d}")
    c_vms = claim(nid(), "ESTIMACIÓN", f"Vaca Muerta: {last.vm_pet_kbbl_d:.0f} mil bbl/d de petróleo "
                  f"({last.vm_share_pet_pct:.0f}% del total) y {last.vm_gas_mmm3_d:.0f} MMm³/d de gas "
                  f"({last.vm_share_gas_pct:.0f}% del total), {lm}",
                  f"{last.vm_share_pet_pct:.0f}% / {last.vm_share_gas_pct:.0f}%", "",
                  cita=f"kbbl/d = m³ × 6,28981 / días / 1000; participación = VM / total; insumos {c_vmp}, {c_vmg}, {c_pt}, {c_gt}")
    c_vmy = claim(nid(), "ESTIMACIÓN", f"Crecimiento interanual del petróleo de Vaca Muerta ({lm} vs. {yago.mes})",
                  f"{(last.vm_pet_m3 / yago.vm_pet_m3 - 1) * 100:.0f}%", "vm",
                  cita=f"{last.vm_pet_m3:.0f} / {yago.vm_pet_m3:.0f} − 1 (formacion='vaca muerta')")
    cell("3 Energía", "Importaciones netas de energía (% del uso)", "ARG", round(imp_ben, 1), 2025, "%", "ESTIMACIÓN",
         c_ben, "menor", "BEN 2025 provisorio (fila adicional; no comparable 1:1 con el WB)")

    # ===================================================================== 4. Litio (USGS MCS 2026)
    li = {"ARG": ("Argentina 513,800 23,000 4,400,000", 13800, 23000, 4_400_000),
          "CHL": ("Chile 548,900 56,000 9,200,000", 48900, 56000, 9_200_000),
          "CAN": ("Canada 4,820 5,600 1,600,000", 4820, 5600, 1_600_000),
          "PRT": ("Portugal 380 380 60,000", 380, 380, 60_000)}
    li_rows, li_cid = [], {}
    for iso, (nd, p24, p25, res) in li.items():
        cid = li_cid[iso] = claim(nid(), "DATO", f"USGS: producción minera de litio de {NOMBRE[iso]} 2024 / 2025e y reservas "
                    "(t de litio contenido; el '5' inicial en ARG/CHL es la nota 'Reported')",
                    f"{p24:,} / {p25:,} t; reservas {res:,} t", "usgs", needle=nd)
        cell("4 Litio", "Reservas de litio (t de Li contenido)", iso, res, 2025, "t", "DATO", cid, "mayor")
        cell("4 Litio", "Producción minera de litio 2025e (t de Li)", iso, p25, 2025, "t", "DATO", cid, "mayor")
        li_rows.append(dict(pais=iso, prod_2024_t=p24, prod_2025e_t=p25, reservas_t=res))
    for iso in ("NZL", "URY"):
        cid = claim(nid(), "DATO", f"{NOMBRE[iso]} no figura en la tabla de producción y reservas de litio del USGS MCS 2026",
                    "no figura", "usgs", needle="Other countries7 — — 2,400,000")
        cell("4 Litio", "Reservas de litio (t de Li contenido)", iso, 0, 2025, "t", "DATO", cid, "mayor",
             "no figura en USGS (se grafica como 0)")
        cell("4 Litio", "Producción minera de litio 2025e (t de Li)", iso, 0, 2025, "t", "DATO", cid, "mayor",
             "no figura en USGS")
        li_rows.append(dict(pais=iso, prod_2024_t=0, prod_2025e_t=0, reservas_t=0))
    c_lw = claim(nid(), "DATO", "USGS: total mundial de producción minera de litio 2024 / 2025e (excl. EE.UU.) y reservas",
                 "222.000 / 290.000 t; reservas 37.000.000 t", "usgs",
                 needle="World total (rounded) 8222,000 8290,000 37,000,000")
    c_lr = claim(nid(), "DATO", "USGS: recursos medidos e indicados de litio de Argentina (los mayores del listado)",
                 "28 millones de t", "usgs", needle="Argentina, 28 million tons; Bolivia, 23 million tons; Chile, 13 million tons")
    c_lsh = claim(nid(), "ESTIMACIÓN", "Participación de Argentina en reservas y producción mundial de litio (2025e)",
                  f"{4.4 / 37 * 100:.1f}% de reservas; {23 / 290 * 100:.1f}% de producción", "usgs",
                  cita=f"4.400.000/37.000.000 y 23.000/290.000; insumos {li_cid['ARG']}, {c_lw}")
    pd.DataFrame(li_rows).to_csv(PROCESSED / "F_litio_usgs.csv", index=False)

    # ===================================================================== 5. Patrimonio Mundial (UNESCO)
    wh = pd.read_csv(FILES["unesco"], sep=";", encoding="utf-8-sig", dtype=str)
    un_rows = []
    for iso in ISO3:
        m = wh[wh.iso_codes.fillna("").str.lower().str.split(",").apply(lambda xs: ISO2[iso] in [x.strip() for x in xs])]
        tot, nat = len(m), int(m.category.isin(["Natural", "Mixed"]).sum())
        cid = claim(nid(), "DATO", f"Sitios del Patrimonio Mundial de {NOMBRE[iso]} (total; naturales + mixtos)",
                    f"{tot}; {nat}", "unesco",
                    cita=f"filas con '{ISO2[iso]}' en iso_codes = {tot}; de ellas category in (Natural, Mixed) = {nat}; "
                         f"dataset modificado 2026-10-02")
        cell("5 Atractivo natural", "Sitios UNESCO Patrimonio Mundial (total)", iso, tot, 2026, "sitios", "DATO", cid, "mayor")
        cell("5 Atractivo natural", "Sitios UNESCO naturales + mixtos", iso, nat, 2026, "sitios", "DATO", cid, "mayor")
        un_rows += [dict(pais=iso, id_no=r.id_no, nombre=r.name_en, categoria=r.category, inscripto=r.date_inscribed,
                         transfronterizo=r.transboundary) for r in m.itertuples()]
    pd.DataFrame(un_rows).to_csv(PROCESSED / "F_unesco_sitios.csv", index=False)

    # ===================================================================== 6. CONTRAPESOS
    # 6a. Defaults soberanos (BoC–BoE)
    wb_ = openpyxl.load_workbook(FILES["boc"], read_only=True)
    rows = list(wb_["Debt_2025"].iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and r[0] == "k")
    db = pd.DataFrame(rows[hdr_i + 1:], columns=rows[hdr_i]).dropna(subset=["k"])
    db = db[db.DEBT_COUNTRY.isin(EN.values())].copy()
    db["DEBT_TOTAL_2025"] = pd.to_numeric(db.DEBT_TOTAL_2025, errors="coerce").fillna(0)
    def_rows = []
    for iso in ISO3:
        s = db[(db.DEBT_COUNTRY == EN[iso]) & (db.DEBT_TOTAL_2025 > 0) & db.DEBT_YEAR.between(1960, 2024)]
        yrs = sorted(int(y) for y in s.DEBT_YEAR)
        if EN[iso] in set(db.DEBT_COUNTRY):
            cid = claim(nid(), "ESTIMACIÓN", f"Años 1960–2024 en que {NOMBRE[iso]} registra deuda soberana en default "
                        "(DEBT_TOTAL > 0) según BoC–BoE", f"{len(yrs)} de 65", "boc",
                        cita=f"hoja Debt_2025; DEBT_COUNTRY={EN[iso]}; conteo de DEBT_YEAR 1960–2024 con DEBT_TOTAL_2025>0; "
                             f"años={','.join(map(str, yrs))}")
            nota = ""
        else:
            cid = claim(nid(), "HIPÓTESIS", f"{NOMBRE[iso]} no figura entre los 166 soberanos/territorios de la base BoC–BoE: "
                        "se interpreta como sin defaults registrados 1960–2024", "0 (no figura)", "boc",
                        cita=f"hoja Debt_2025: ningún registro con DEBT_COUNTRY={EN[iso]}")
            nota = "no figura en la base"
        cell("6 Contrapesos", "Años con deuda soberana en default, 1960–2024 (de 65)", iso, len(yrs), "1960–2024",
             "años", "ESTIMACIÓN" if not nota else "HIPÓTESIS", cid, "menor", nota)
        def_rows += [dict(pais=iso, anio=y, en_default=int(y in yrs)) for y in range(1960, 2025)]
    pd.DataFrame(def_rows).to_csv(PROCESSED / "F_defaults_anios.csv", index=False)
    for y in (2001, 2020):
        r = db[(db.DEBT_COUNTRY == "Argentina") & (db.DEBT_YEAR == y)].iloc[0]
        claim(nid(), "DATO", f"BoC–BoE: deuda de Argentina en default en {y}, bonos en moneda extranjera",
              f"USD {r.DEBT_FC_BONDS_2025:,.0f} millones", "boc",
              cita=f"hoja Debt_2025; k=ARG_{y}; DEBT_FC_BONDS_2025={r.DEBT_FC_BONDS_2025}; DEBT_TOTAL_2025={r.DEBT_TOTAL_2025}")
    r = db[(db.DEBT_COUNTRY == "Argentina") & (db.DEBT_YEAR == 2024)].iloc[0]
    claim(nid(), "DATO", "BoC–BoE: deuda de Argentina todavía clasificada en default en 2024", f"USD {r.DEBT_TOTAL_2025:,.0f} millones",
          "boc", cita=f"hoja Debt_2025; k=ARG_2024; DEBT_TOTAL_2025={r.DEBT_TOTAL_2025}")
    claim(nid(), "DATO", "Metodología BoC–BoE: las reestructuraciones de préstamos oficiales de la UE a Grecia, Irlanda y "
          "Portugal (2013, 2018) cuentan como default por pérdida en valor presente (explica el registro de Portugal 2013)",
          "Portugal 2013", "boc_note",
          needle="These official debt restructurings are consistent with our definition of sovereign defaults because "
                 "they result in creditor losses in present-value terms")

    # 6b. Inflación
    cpi = load_d360("wdi_cpi")
    infl_rows = []
    for iso in ISO3:
        s = cpi[(cpi.REF_AREA == iso) & cpi.OBS_VALUE.notna()].sort_values("TIME_PERIOD")
        for y in (2024, 2025):
            v = s[s.TIME_PERIOD == y]
            if len(v):
                val = v.OBS_VALUE.iloc[0]
                cid = claim(nid(), "DATO", f"Inflación (IPC, % anual promedio) de {NOMBRE[iso]}, {y} — Banco Mundial",
                            f"{val:.1f}%", "wdi_cpi",
                            cita=f"INDICATOR=WB_WDI_FP_CPI_TOTL_ZG; REF_AREA={iso}; TIME_PERIOD={y}; OBS_VALUE={val}")
                cell("6 Contrapesos", f"Inflación {y} (IPC, % promedio anual)", iso, round(val, 1), y, "%", "DATO", cid, "menor")
        infl_rows += [dict(pais=iso, anio=int(r.TIME_PERIOD), inflacion_pct=r.OBS_VALUE) for r in s.itertuples()]
    pd.DataFrame(infl_rows).to_csv(PROCESSED / "F_inflacion_wb.csv", index=False)

    ipc = pd.read_csv(FILES["ipc"])
    ipc.columns = ["mes", "ipc_nivel", "var_mensual", "var_interanual"]
    ipc["mes"] = ipc.mes.str[:7]
    ipc.to_csv(PROCESSED / "F_ipc_arg_mensual.csv", index=False)
    lt = ipc.iloc[-1]
    c_ipcm = claim(nid(), "DATO", f"IPC Nacional INDEC: variación mensual y interanual, {lt.mes}",
                   f"{lt.var_mensual * 100:.1f}% m/m; {lt.var_interanual * 100:.1f}% i.a.", "ipc",
                   cita=f"indice_tiempo={lt.mes}-01; var_pct={lt.var_mensual}; var_pct_ia={lt.var_interanual}")
    dec = ipc[ipc.mes == "2025-12"].iloc[0]
    claim(nid(), "DATO", "IPC Nacional INDEC: inflación diciembre 2025 / diciembre 2024", f"{dec.var_interanual * 100:.1f}%",
          "ipc", cita=f"indice_tiempo=2025-12-01; var_pct_ia={dec.var_interanual}")
    yr = ipc.assign(a=ipc.mes.str[:4]).groupby("a").ipc_nivel.mean()
    avg25 = (yr["2025"] / yr["2024"] - 1) * 100
    c_a25 = claim(nid(), "ESTIMACIÓN", "Argentina 2025: inflación promedio anual (misma definición que FP.CPI.TOTL.ZG) "
                  "a partir del IPC INDEC (el Banco Mundial aún no publica 2025 para Argentina)", f"{avg25:.1f}%", "ipc",
                  cita=f"promedio(IPC 2025)/promedio(IPC 2024) − 1 = {yr['2025']:.2f}/{yr['2024']:.2f} − 1")
    cell("6 Contrapesos", "Inflación 2025 (IPC, % promedio anual)", "ARG", round(avg25, 1), 2025, "%", "ESTIMACIÓN", c_a25,
         "menor", "INDEC vía datos.gob.ar; WB sin dato 2025")

    # 6c. Riesgo país y calificaciones: ver sección 8 (claims F123 en adelante, para no renumerar los existentes)

    # 6d. Controles de capital y corralito (normativa)
    cc = [
        (nid(), "DATO", "Corralito: el Decreto 1570/2001 limitó los retiros en efectivo a $250 o USD 250 por semana",
         "01/12/2001", "dec1570",
         "Los retiros en efectivo que superen los PESOS DOSCIENTOS CINCUENTA ($ 250) o DOLARES ESTADOUNIDENSES "
         "DOSCIENTOS CINCUENTA (U$S 250) por semana"),
        (nid(), "DATO", "Corralito: el Decreto 1570/2001 prohibió las transferencias al exterior (salvo excepciones)",
         "01/12/2001", "dec1570", "b) Las transferencias al exterior, con excepción de las que correspondan a operaciones de comercio exterior"),
        (nid(), "DATO", "Cepo 2011: validación fiscal (AFIP) obligatoria para comprar moneda extranjera, vigente desde el 11/11/2011",
         "11/11/2011", "a5245", "se ha dispuesto con vigencia a partir del 11.11.2011 inclusive"),
        (nid(), "DATO", "Salida del cepo 2015: la Com. A 5850 deja sin efecto la consulta AFIP, con vigencia desde el 17/12/2015",
         "17/12/2015", "a5850", "con vigencia a partir del 17.12.2015 inclusive, se ha dispuesto lo siguiente: 1. Dejar sin "
                                "efecto la consulta y registro de las operaciones cambiarias"),
        (nid(), "DATO", "Cepo 2019: el DNU 609/2019 faculta al BCRA a exigir autorización previa para comprar divisas",
         "01/09/2019", "dnu609", "Decreto 609/2019 DNU-2019-609-APN-PTE Ciudad de Buenos Aires, 01/09/2019"),
        (nid(), "DATO", "Cepo 2019: la Com. A 6770 exige conformidad previa del BCRA a personas humanas por encima de USD 10.000 mensuales",
         "01/09/2019", "a6770", "cuando supere el equivalente de US$ 10.000 mensuales"),
        (nid(), "DATO", "Cepo 2019: la Com. A 6815 limita a USD 200 mensuales la compra de personas humanas sin conformidad previa",
         "28/10/2019", "a6815", "cuando su- pere el equivalente de US$ 200 mensuales"),
        (nid(), "DATO", "Salida del cepo 2025 (personas humanas): la Com. A 8226 rige desde el 14/04/2025 y habilita la compra "
         "sin conformidad previa del BCRA", "14/04/2025", "a8226",
         "adoptó la siguiente resolución con vigencia a partir del 14/04/25: 1. Establecer que las entidades podrán dar acceso "
         "al mercado de cambios a las personas humanas residentes, sin conformidad previa del Banco Central"),
        (nid(), "DATO", "La Com. A 8226 mantiene restricciones para empresas: solo habilita girar dividendos de ejercicios "
         "iniciados desde el 01/01/2025", "01/01/2025", "a8226",
         "realizadas en estados contables anuales regulares y auditados de ejercicios iniciados a partir del 01/01/25"),
    ]
    for cid, lab, af, val, sid, nd in cc:
        claim(cid, lab, af, val, sid, needle=nd)
    dias = ((pd.Timestamp("2015-12-17") - pd.Timestamp("2011-11-11")).days
            + (pd.Timestamp("2025-04-14") - pd.Timestamp("2019-09-01")).days) / 365.25
    c_cepo = claim(nid(), "ESTIMACIÓN", "Años con cepo cambiario para personas humanas entre 11/2011 y 04/2025",
                   f"{dias:.1f} años", "", cita=f"(17/12/2015 − 11/11/2011) + (14/04/2025 − 01/09/2019); insumos "
                                                f"{cc[2][0]}, {cc[3][0]}, {cc[4][0]}, {cc[7][0]}")
    cell("6 Contrapesos", "Años con cepo cambiario (personas), 2011–2025", "ARG", round(dias, 1), "2011–2025", "años",
         "ESTIMACIÓN", c_cepo, "menor")
    for iso in ISO3[1:]:
        cell("6 Contrapesos", "Años con cepo cambiario (personas), 2011–2025", iso, None, "", "años", "n/r", "", "menor",
             "no relevado en este módulo")

    # 6e. WJP Rule of Law Index 2025
    wjp = {"ARG": ("Argentina 0.54 -1.0% 65 -1", 0.54, 65), "NZL": ("New Zealand 0.83 0.0% 5 1", 0.83, 5),
           "URY": ("Uruguay 0.72 -0.2% 23 1", 0.72, 23), "CHL": ("Chile 0.66 -0.3% 35 1", 0.66, 35),
           "PRT": ("Portugal 0.67 -0.5% 29 -1", 0.67, 29), "CAN": ("Canada 0.79 -0.4% 13 -1", 0.79, 13)}
    for iso, (nd, sc, rk) in wjp.items():
        cid = claim(nid(), "DATO", f"WJP Rule of Law Index 2025: puntaje (0–1), variación, puesto (de 143) y cambio de puesto "
                    f"de {NOMBRE[iso]}", f"{sc} / puesto {rk}", "wjp", needle=nd)
        cell("6 Contrapesos", "WJP Rule of Law Index (0–1)", iso, sc, 2025, "0 a 1", "DATO", cid, "mayor", f"puesto {rk} de 143")
    claim(nid(), "DATO", "El WJP Rule of Law Index 2025 cubre 143 países y jurisdicciones", "143", "wjp",
          needle="scholars across 143 countries and jurisdictions")

    # ===================================================================== 7. Geopolítica (solo hechos con fuente primaria)
    c_geo1 = claim(nid(), "DATO", "Argentina y EE.UU. suscribieron el Acuerdo sobre Comercio e Inversiones Recíprocos (05/02/2026)",
                   "05/02/2026", "cancilleria",
                   needle="La República Argentina y los Estados Unidos suscribieron el Acuerdo sobre Comercio e Inversiones Recíprocos")
    c_geo2 = claim(nid(), "DATO", "El Tesoro de EE.UU. (ESF) firmó un acuerdo de estabilización cambiaria por USD 20.000 millones con el "
                   "BCRA y en octubre de 2025 ejecutó un swap por USD 2.500 millones", "USD 20.000 M / 2.500 M", "esf",
                   needle="The ESF also has an exchange stabilization agreement (ESA) with the Central Bank of Argentina (BCRA) for $20 billion")
    claim(nid(), "DATO", "Detalle del swap ESF–BCRA de octubre de 2025", "USD 2.500 M", "esf",
          needle="the BCRA exchanged pesos for $2.5 billion")

    # ===================================================================== 8. Pendientes cerrados (2026-10-03)
    # 8a. Riesgo país de Argentina: no hay serie pública del EMBI, pero el BCRA lo cita textualmente en el IPOM
    c_rp = claim(nid(), "DATO", "Riesgo país de Argentina (spread EMBI, citado por el BCRA): bajó de 556 a 434 p.b. en los tres "
                 "meses hasta fines de julio de 2026", "434 p.b. (fines de jul-2026)", "ipom",
                 needle="nuyó 122 p.b. en los últimos 3 meses hasta fines de julio (de 556 p.b. a 434 p.b.)")
    c_rpb = claim(nid(), "DATO", "Prima de riesgo soberano promedio de los países con calificación B- (BCRA, IPOM)", "322 p.b.",
                  "ipom", needle="aproximándose así a la prima de riesgo soberano que exhiben en promedio el resto de los "
                                 "créditos con similar calificación crediticia (322 p.b.)")
    c_rpl = claim(nid(), "DATO", "EMBI Latam (referencia regional, BCRA): en torno a 260 p.b. a fines de julio de 2026",
                  "≈260 p.b.", "ipom",
                  needle="el EMBI Latam se ubicó en torno a 260 puntos básicos (p.b.) a fines de julio")
    claim(nid(), "DATO", "El riesgo país de Argentina alcanzó el nivel más bajo desde principios de 2018 (BCRA, IPOM 2T-2026)",
          "mínimo desde 2018", "ipom", needle="riesgo país, que alcanzó el nivel más bajo desde principios de 2018")
    cell("6 Contrapesos", "Riesgo país: spread EMBI (p.b.)", "ARG", 434, "jul-2026", "p.b.", "DATO", c_rp, "menor",
         f"BCRA IPOM 2T-2026; promedio de soberanos B-: 322 p.b. ({c_rpb}); EMBI Latam ≈260 p.b. ({c_rpl})")
    for iso in ISO3[1:]:
        cell("6 Contrapesos", "Riesgo país: spread EMBI (p.b.)", iso, None, "", "p.b.", "s/d", "", "menor",
             "sin fuente pública/oficial con el dato por país (EMBI de JP Morgan es propietario)")

    # 8b. Calificación soberana de largo plazo en moneda extranjera (S&P / Moody's / Fitch), según cada gobierno
    rating_src = {
        "ARG": ("ipom", ("B-", "B3", "B-"), "jul-2026",
                ["la agencia Fitch elevó la calificación de la deuda soberana argentina, pasando de CCC+ a B-, y "
                 "aproximadamente un mes después la agencia Standard & Poor’s hizo lo propio (de CCC a B-)",
                 "la agencia Moody’s llevó la calificación crediticia de Caa1 a B3"]),
        "NZL": ("rt_nz", ("AA+", "Aaa", "AA+"), "abr-2026",
                ["Moody's Investors Service Aaa (stable outlook) Aaa (stable outlook) 16 April 2024 S&P Global Ratings "
                 "AAA (stable outlook) AA+ (stable outlook) 14 October 2025 Fitch Ratings AA+ (negative outlook) "
                 "AA+ (negative outlook) 20 March 2026"]),
        "URY": ("rt_uy", ("BBB+", "Baa1", "BBB"), "sep-2026",
                ["Fitch Ratings BBB BBB Estable Setiembre 2026", "S&P BBB+ BBB+ Estable Noviembre 2025",
                 "Baa1 Baa1 Estable Marzo 2024"]),
        "CHL": ("rt_cl", ("A", "A2", "A-"), "última entrada: oct-2024",
                ["16-10-2024 A Estable", "15-09-2022 A2 Estable", "15-10-2020 A− Estable"]),
        "PRT": ("rt_pt", ("A+", "A3", "A+"), "sep-2026",
                ["currently at A3 |Stable byMoody’s; A+ |Positive by S&P; A+ |Stableby Fitch"]),
        "CAN": ("rt_ca", ("AAA", "Aaa", "AA+"), "31/03/2025",
                ["These strengths are reflected in Canada's strong current credit ratings: Moody's (Aaa), S&P (AAA), "
                 "Fitch (AA+), DBRS (AAA)."]),
    }
    rt_rows = []
    for iso, (sid, (sp, md, fi), fecha, nds) in rating_src.items():
        cid = claim(nid(), "DATO", f"Calificación soberana de largo plazo en moneda extranjera de {NOMBRE[iso]} "
                    "(S&P / Moody's / Fitch), según fuente oficial", f"{sp} / {md} / {fi} ({fecha})", sid, needles=nds)
        notches = [NOTCH_SP.index(sp), NOTCH_MD.index(md), NOTCH_SP.index(fi)]
        avg = sum(notches) / 3
        cid_e = claim(nid(), "ESTIMACIÓN", f"{NOMBRE[iso]}: escalones por debajo de AAA/Aaa, promedio de S&P, Moody's y Fitch "
                      "(BBB-/Baa3 = 9 = último escalón de grado de inversión)", f"{avg:.1f}", "",
                      cita=f"AAA/Aaa=0, AA+/Aa1=1, …, BBB-/Baa3=9, …, B-/B3=15; ({notches[0]}+{notches[1]}+{notches[2]})/3; "
                           f"insumo {cid}")
        cell("6 Contrapesos", "Calificación soberana LP m/e (S&P / Moody's / Fitch)", iso, f"{sp} / {md} / {fi}", fecha,
             "letras", "DATO", cid, "mayor")
        cell("6 Contrapesos", "Calificación soberana: escalones bajo AAA (prom. 3 agencias)", iso, round(avg, 1), fecha,
             "escalones", "ESTIMACIÓN", cid_e, "menor", "grado de inversión ≤ 9")
        rt_rows.append(dict(pais=iso, sp=sp, moodys=md, fitch=fi, fecha=fecha, escalones_sp=notches[0],
                            escalones_moodys=notches[1], escalones_fitch=notches[2], escalones_prom=round(avg, 2),
                            grado_inversion=avg <= 9, fuente=SRC[sid][3], claim_dato=cid, claim_estimacion=cid_e))
    pd.DataFrame(rt_rows).to_csv(PROCESSED / "F_calificaciones_soberanas.csv", index=False)

    # 8c. Rendimiento de bonos del gobierno a largo plazo (OCDE). Argentina y Uruguay no son miembros: sin serie.
    lt = pd.read_csv(FILES["oecd_lt"], dtype=str)
    lt["OBS_VALUE_f"] = pd.to_numeric(lt.OBS_VALUE, errors="coerce")
    (lt[["REF_AREA", "TIME_PERIOD", "OBS_VALUE_f"]].rename(columns={"REF_AREA": "pais", "TIME_PERIOD": "mes",
                                                                    "OBS_VALUE_f": "tasa_largo_plazo_pct"})
     .sort_values(["pais", "mes"]).to_csv(PROCESSED / "F_oecd_tasa_largo_plazo.csv", index=False))
    for iso in ISO3:
        s = lt[(lt.REF_AREA == iso) & lt.OBS_VALUE_f.notna()].sort_values("TIME_PERIOD")
        if len(s):
            r = s.iloc[-1]
            cid = claim(nid(), "DATO", f"OCDE: tasa de interés de largo plazo (bonos del gobierno ~10 años, moneda local) de "
                        f"{NOMBRE[iso]}, {r.TIME_PERIOD}", f"{r.OBS_VALUE_f:.2f}% anual", "oecd_lt",
                        cita=f"REF_AREA={iso}; MEASURE=IRLT; TIME_PERIOD={r.TIME_PERIOD}; OBS_VALUE={r.OBS_VALUE}")
            cell("6 Contrapesos", "Tasa de largo plazo, bonos del gobierno ~10 años en moneda local (% anual, OCDE)", iso,
                 round(r.OBS_VALUE_f, 2), r.TIME_PERIOD, "%", "DATO", cid, "menor", "no es un spread: incluye inflación esperada")
        else:
            cell("6 Contrapesos", "Tasa de largo plazo, bonos del gobierno ~10 años en moneda local (% anual, OCDE)", iso,
                 None, "", "%", "s/d", "", "menor", "no es miembro de la OCDE: la consulta no devuelve serie")

    # 8d. Controles de capital vigentes (texto ordenado al 14/09/2026; no hay circulares CAMEX posteriores hasta la A 8488)
    c_to = claim(nid(), "DATO", "Texto ordenado de Exterior y Cambios vigente: al 14/09/2026, última comunicación incorporada A 8481",
                 "14/09/2026", "texord", needle="-Última comunicación incorporada: A 8481- Texto ordenado al 14/09/2026")
    c_pj = claim(nid(), "DATO", "Empresas (personas jurídicas): siguen necesitando conformidad previa del BCRA para formar activos "
                 "externos (texto ordenado, punto 3.10)", "vigente al 14/09/2026", "texord",
                 needle="El acceso al mercado de cambios por parte de personas jurídicas que no sean entidades autorizadas a operar "
                        "en cambios, gobiernos locales, Fondos Comunes de Inversión, Fideicomisos y otras universalidades "
                        "constituidas en el país, requerirá la conformidad previa del BCRA para la formación de activos externos")
    c_div = claim(nid(), "DATO", "Dividendos: el acceso sigue limitado a utilidades de ejercicios iniciados desde el 01/01/2025 "
                  "(salvo excepciones: BOPREAL, RIGI, aportes desde 17/01/2020, etc.; punto 3.4.4)", "vigente al 14/09/2026",
                  "texord", needle="se trata de utilidades distribuibles obtenidas a partir de ganancias realizadas en estados "
                                   "contables anuales regulares y auditados de ejercicios iniciados a partir del 01/01/25")
    c_ph = claim(nid(), "DATO", "Personas humanas: acceso sin límite para comprar billetes o depósitos en moneda extranjera con "
                 "débito en cuenta; las personas jurídicas requieren conformidad previa (BCRA)", "sin límite (personas)",
                 "bcra_nec", needles=[
                     "Las personas humanas residentes pueden acceder sin límite al mercado de cambios para formar activos "
                     "externos en billetes o depósitos, siempre que la operación se realice mediante débito en cuenta en una "
                     "entidad financiera local",
                     "Las personas jurídicas, requieren la conformidad previa del BCRA para acceder al mercado de cambios con "
                     "esos fines",
                     "Las personas no residentes también deben contar con la conformidad previa del BCRA"])
    c_90 = claim(nid(), "DATO", "Personas humanas: quien compra divisas se compromete a no operar títulos con liquidación en "
                 "moneda extranjera por 90 días (restricción cruzada, punto 3.8.5); otras modalidades de formación de activos "
                 "externos, tope de USD 200 mensuales (punto 3.9.1)", "90 días; USD 200", "texord", needles=[
                     "compras de títulos valores con liquidación en moneda extranjera a partir del momento en que requiere el "
                     "acceso y por los 90 (noventa) días corridos subsiguientes",
                     "3.9.1. El cliente no supere, en el mes calendario en el conjunto de las entidades y por el conjunto de los "
                     "conceptos señalados, el equivalente a USD 200"])
    c_dvu = claim(nid(), "DATO", "La flexibilización cambiaria permitió a las empresas, tras seis años, girar dividendos por "
                  "~USD 2.800 millones en el primer semestre de 2026 (BCRA, IPOM)", "USD 2.800 M (1S-2026)", "ipom",
                  needle="permitió a las empresas, luego de seis años, comenzar a girar dividendos por aproximadamente "
                         "USD 2.800 millones en el primer semestre del año")

    # 8e. Geopolítica: estatus de aliado extra-OTAN (State Department, vía Wayback)
    c_mnna = claim(nid(), "DATO", "Argentina (y Nueva Zelanda) figuran entre los 19 países designados por EE.UU. como "
                   "Major Non-NATO Ally", "19 MNNA; incluye ARG y NZL", "mnna",
                   needle="Currently 19 countries are designated as MNNAs under 22 U.S.C. §2321k and 10 U.S.C. §2350a "
                          "Argentina, Australia, Bahrain, Brazil, Colombia, Egypt, Israel, Japan, Jordan, Kenya, Kuwait, "
                          "Morocco, New Zealand")

    # ===================================================================== tabla resumen + gráficos
    tab = pd.DataFrame(table)
    tab = tab.sort_values("dimension", kind="stable").reset_index(drop=True)
    tab.to_csv(PROCESSED / "F_tabla_resumen.csv", index=False)
    wide = tab.pivot_table(index=["dimension", "indicador"], columns="pais", values="valor", aggfunc="first")[ISO3]
    wide.to_csv(PROCESSED / "F_tabla_resumen_ancha.csv")

    charts(tab, wgi_est, en, ipc, pd.DataFrame(def_rows))
    write_ledger(MODULO, CLAIMS)
    print(f"{len(CLAIMS)} afirmaciones -> docs/claims/claims_F.csv; {len(tab)} celdas en F_tabla_resumen.csv")
    for k, v in dict(gpi_n=gpi_n, gpi_arg=gpi_arg, ben=c_ben, vm=c_vms, vmy=c_vmy, lsh=c_lsh, lr=c_lr, ipc=c_ipcm,
                     a25=c_a25, cepo=c_cepo, geo1=c_geo1, geo2=c_geo2, rp=c_rp, rpb=c_rpb, rpl=c_rpl, to=c_to,
                     pj=c_pj, div=c_div, ph=c_ph, d90=c_90, dvu=c_dvu, mnna=c_mnna).items():
        c = next(x for x in CLAIMS if x["claim_id"] == v)
        print(f"  {v} [{k}] {c['afirmacion'][:90]} = {c['valor']}")


# ------------------------------------------------------------------------------------------------ gráficos
ACCENT, GRAY, INK, INK2, GRID = "#2a78d6", "#898781", "#0b0b0b", "#52514e", "#e1e0d9"
SOURCE_FS = 7


def _style(ax):
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.spines["bottom"].set_color("#c3c2b7")
    ax.tick_params(colors=INK2, labelsize=8, length=0)
    ax.grid(axis="x", color=GRID, linewidth=0.6)
    ax.set_axisbelow(True)


def _save(fig, name, fuente):
    fig.text(0.01, 0.005, f"Fuente: {fuente}. Elaboración: Colossus Lab.", fontsize=SOURCE_FS, color=INK2, ha="left", va="bottom")
    for ext in ("png", "svg"):
        fig.savefig(CHARTS / f"F_{name}.{ext}", dpi=150, facecolor="white")
    plt.close(fig)


def charts(tab, wgi_est, en, ipc, dfl):
    plt.rcParams.update({"font.family": "DejaVu Sans", "figure.facecolor": "white", "axes.facecolor": "white"})
    # --- 1) Small multiples: un panel por indicador, filas = países, Argentina resaltada
    panels = [
        ("WGI Rule of Law (estimación)", "WGI Rule of Law (estimación)", "mayor"),
        ("WGI Political Stability and Absence of Violence (estimación)", "WGI Estabilidad política", "mayor"),
        ("WGI Government Effectiveness (estimación)", "WGI Efectividad del gobierno", "mayor"),
        ("WGI Control of Corruption (estimación)", "WGI Control de la corrupción", "mayor"),
        ("Global Peace Index: puntaje (menor = más pacífico)", "Global Peace Index 2026 (puntaje)", "menor"),
        ("Autosuficiencia en cereales (prod./suministro interno)", "Autosuficiencia en cereales, % (2019–23)", "mayor"),
        ("Cobertura calórica doméstica (kcal ponderadas, SSR tope 100%)", "Cobertura calórica doméstica, % (2019–23)", "mayor"),
        ("Importaciones netas de energía (% del uso)", "Import. netas de energía, % del uso (WB 2023)", "menor"),
        ("Reservas de litio (t de Li contenido)", "Reservas de litio, millones de t (USGS)", "mayor"),
        ("Sitios UNESCO naturales + mixtos", "Sitios UNESCO naturales + mixtos (2026)", "mayor"),
        ("Inflación 2024 (IPC, % promedio anual)", "Inflación 2024, % (WB)", "menor"),
        ("Años con deuda soberana en default, 1960–2024 (de 65)", "Años en default 1960–2024 (BoC–BoE)", "menor"),
        ("WJP Rule of Law Index (0–1)", "WJP Rule of Law Index 2025 (0–1)", "mayor"),
        ("Calificación soberana: escalones bajo AAA (prom. 3 agencias)",
         "Calificación soberana: escalones bajo AAA\n(prom. S&P, Moody's, Fitch; grado de inversión ≤ 9)", "menor"),
    ]
    order = ISO3[::-1]
    fig, axes = plt.subplots(5, 3, figsize=(12, 13.5))
    axes = axes.ravel()
    for ax, (ind, title, better) in zip(axes, panels):
        d = tab[(tab.indicador == ind) & (tab.etiqueta.isin(["DATO", "ESTIMACIÓN", "HIPÓTESIS"]))]
        if ind.startswith("Importaciones"):
            d = d[d.etiqueta == "DATO"]
        d = d.drop_duplicates("pais").set_index("pais")
        vals = d.valor.astype(float)
        if ind.startswith("Reservas"):
            vals = vals / 1e6
        ys = np.arange(len(order))
        lo = min(0, vals.min())
        for y, iso in zip(ys, order):
            if iso not in vals:
                continue
            v = vals[iso]
            c = ACCENT if iso == "ARG" else GRAY
            ax.plot([lo, v], [y, y], color=GRID, linewidth=1, zorder=1)
            ax.scatter([v], [y], s=46 if iso == "ARG" else 30, color=c, zorder=3, edgecolor="white", linewidth=1.5)
            dec = 2 if (vals.abs().max() < 10 and not (vals == vals.round()).all()) else 0
            dec = 1 if (dec == 0 and 0 < abs(v) < 1) else dec
            dec = 1 if ind.startswith("Calificación") else dec
            lab = f"{v:,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")
            if ind.startswith("Reservas") and v == 0:
                lab = "no figura"
            ax.annotate(lab, (v, y), xytext=(6, 0),
                        textcoords="offset points", va="center", fontsize=7.5,
                        color=INK if iso == "ARG" else INK2, fontweight="bold" if iso == "ARG" else "normal")
        ax.set_yticks(ys, [NOMBRE[i] for i in order])
        for t in ax.get_yticklabels():
            if t.get_text() == "Argentina":
                t.set_fontweight("bold"); t.set_color(INK)
        rng = vals.max() - lo
        ax.set_xlim(lo - 0.05 * rng, vals.max() + 0.28 * rng)
        if lo < 0 < vals.max():
            ax.axvline(0, color="#c3c2b7", linewidth=0.8)
        if ind.startswith("Calificación"):  # frontera del grado de inversión (BBB-/Baa3)
            ax.axvline(9, color="#c3c2b7", linewidth=0.8, linestyle="--")
        _style(ax)
        ax.set_title(f"{title}\n{'mayor' if better == 'mayor' else 'menor'} = mejor para el refugio",
                     fontsize=8.5, color=INK, loc="left")
    for ax in axes[len(panels):]:
        ax.axis("off")
    rp = tab[(tab.indicador.str.startswith("Riesgo país")) & (tab.pais == "ARG")].iloc[0]
    inf25 = f"{tab[(tab.indicador.str.startswith('Inflación 2025')) & (tab.pais == 'ARG')].valor.iloc[0]:.1f}".replace(".", ",")
    axes[-1].text(0, 0.9, f"Riesgo país Argentina (BCRA): {int(rp.valor)} p.b.\n({rp.anio}); promedio de soberanos B-: 322 p.b.\n"
                  "Sin dato público por país para los otros cinco\n(el EMBI de JP Morgan es propietario).\n\n"
                  "Inflación 2025 Argentina (INDEC, ESTIMACIÓN):\n"
                  f"{inf25}% promedio anual.\n\n"
                  "Sin índice compuesto: cada panel es una dimensión.", fontsize=8.5, color=INK2, va="top",
                  transform=axes[-1].transAxes)
    fig.suptitle("Argentina se destaca en alimentos, litio y energía, y queda última de las seis en paz, instituciones y estabilidad macro",
                 fontsize=12, color=INK, x=0.01, ha="left", y=0.995)
    fig.tight_layout(rect=(0, 0.02, 1, 0.985))
    _save(fig, "small_multiples", "Banco Mundial (WGI, WDI), IEP (GPI 2026), FAOSTAT FBS, USGS MCS 2026, UNESCO (whc001), "
          "BoC–BoE 2025, WJP 2025, BCRA (IPOM 2T-2026), oficinas de deuda (calificaciones)")

    # --- 2) WGI Rule of Law: serie 1996–2025
    fig, ax = plt.subplots(figsize=(9, 5))
    s = wgi_est[wgi_est.indicador == "Rule of Law"]
    for iso in ISO3[1:] + ["ARG"]:
        d = s[s.pais == iso].sort_values("anio")
        c, lw = (ACCENT, 2) if iso == "ARG" else (GRAY, 1.2)
        ax.plot(d.anio, d.estimacion, color=c, linewidth=lw)
        ax.annotate(NOMBRE[iso], (d.anio.iloc[-1], d.estimacion.iloc[-1]), xytext=(5, 0), textcoords="offset points",
                    va="center", fontsize=8, color=INK if iso == "ARG" else INK2)
    ax.axhline(0, color="#c3c2b7", linewidth=0.8)
    _style(ax); ax.grid(axis="y", color=GRID, linewidth=0.6); ax.grid(axis="x", visible=False)
    ax.set_xlim(1995, 2029); ax.set_ylabel("Estimación (−2,5 a +2,5)", fontsize=8, color=INK2)
    ax.set_title("Estado de derecho (WGI): Argentina, única de las seis bajo cero desde 2000; mejora desde 2023",
                 fontsize=11, color=INK, loc="left")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    _save(fig, "wgi_rule_of_law", "Banco Mundial, Worldwide Governance Indicators (API Data360)")

    # --- 3) Vaca Muerta: petróleo y gas (dos paneles, una escala por panel)
    d = en.dropna(subset=["vm_pet_m3"]).copy()
    d["fecha"] = pd.to_datetime(d.mes + "-01")
    d = d[d.fecha >= "2015-01-01"]
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(11, 4.5))
    for ax, tot, vmc, unit, tt in ((a1, "pet_total_kbbl_d", "vm_pet_kbbl_d", "miles de barriles por día", "Petróleo"),
                                   (a2, "gas_total_mmm3_d", "vm_gas_mmm3_d", "millones de m³ por día", "Gas natural")):
        ax.plot(d.fecha, d[tot], color=GRAY, linewidth=1.2, label="Total país")
        ax.plot(d.fecha, d[vmc], color=ACCENT, linewidth=2, label="Vaca Muerta")
        ax.legend(loc="upper left", frameon=False, fontsize=8)
        _style(ax); ax.grid(axis="y", color=GRID, linewidth=0.6); ax.grid(axis="x", visible=False)
        ax.set_ylim(0, None); ax.set_title(f"{tt} ({unit})", fontsize=9.5, color=INK, loc="left")
    fig.suptitle(f"Vaca Muerta ya aporta {d.vm_share_pet_pct.iloc[-1]:.0f}% del petróleo y {d.vm_share_gas_pct.iloc[-1]:.0f}% "
                 f"del gas de Argentina ({d.mes.iloc[-1]}, dato provisorio)", fontsize=11, color=INK, x=0.01, ha="left")
    fig.tight_layout(rect=(0, 0.04, 1, 0.97))
    _save(fig, "vaca_muerta", "Secretaría de Energía, Capítulo IV (producción por pozo no convencional y series por cuenca)")

    # --- 4) Inflación mensual Argentina (INDEC)
    d = ipc.dropna(subset=["var_mensual"]).copy()
    d["fecha"] = pd.to_datetime(d.mes + "-01")
    fig, ax = plt.subplots(figsize=(10, 4.5))
    ax.bar(d.fecha, d.var_mensual * 100, width=24, color=ACCENT)
    lt = d.iloc[-1]
    mx = d.loc[d.var_mensual.idxmax()]
    for r, off in ((mx, 6), (lt, 22)):
        ax.annotate(f"{r.mes}: {r.var_mensual * 100:.1f}%".replace(".", ","), (r.fecha, r.var_mensual * 100),
                    xytext=(0, off), textcoords="offset points", ha="center", fontsize=8, color=INK)
    _style(ax); ax.grid(axis="y", color=GRID, linewidth=0.6); ax.grid(axis="x", visible=False)
    ax.set_ylabel("% mensual", fontsize=8, color=INK2)
    ax.set_title(f"Inflación mensual de Argentina: bajó desde el pico de {mx.var_mensual * 100:.1f}% ({mx.mes}) pero sigue en "
                 f"{lt.var_interanual * 100:.0f}% interanual ({lt.mes})".replace(".", ","), fontsize=10.5, color=INK, loc="left")
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    _save(fig, "inflacion_arg", "INDEC, IPC Nacional base dic-2016, vía API de Series de Tiempo (datos.gob.ar)")

    # --- 5) Defaults: años con deuda en default (tira por país)
    fig, ax = plt.subplots(figsize=(10, 3.4))
    for y, iso in enumerate(order):
        yrs = dfl[(dfl.pais == iso) & (dfl.en_default == 1)].anio
        ax.barh([y] * len(yrs), [0.9] * len(yrs), left=yrs - 0.45, height=0.6, color=ACCENT if iso == "ARG" else GRAY)
        ax.text(2025.5, y, f"{len(yrs)} año" + ("" if len(yrs) == 1 else "s"), va="center", fontsize=8, color=INK if iso == "ARG" else INK2)
    ax.set_yticks(range(len(order)), [NOMBRE[i] for i in order])
    ax.set_xlim(1959, 2031)
    _style(ax)
    ax.set_title("Años con deuda soberana en default, 1960–2024: Argentina suma más que los otros cinco juntos",
                 fontsize=10.5, color=INK, loc="left")
    fig.tight_layout(rect=(0, 0.06, 1, 1))
    _save(fig, "defaults", "Bank of Canada–Bank of England Sovereign Default Database 2025 (NZ y Canadá no figuran en la base)")


if __name__ == "__main__":
    main()
